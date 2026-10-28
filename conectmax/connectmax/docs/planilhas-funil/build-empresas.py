#!/usr/bin/env python3
"""Gera empresas.json a partir do portal (Contas a receber) + planilhas Pedidos/Orçamentos.

Uso:
  # 1) Login + dump CR (requer ~/.config/conectmax-franqueado/credentials.env)
  python3 build-empresas.py --fetch-portal

  # 2) Só regenerar a partir de /tmp/cr-all.html (ou --html PATH) + xlsx locais
  python3 build-empresas.py

Saída: empresas.json nesta pasta.
"""
from __future__ import annotations

import argparse
import json
import re
import subprocess
import sys
from collections import Counter, defaultdict
from datetime import date
from html import unescape
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
COOKIE = Path.home() / ".config/conectmax-franqueado/cookies.txt"
CREDS = Path.home() / ".config/conectmax-franqueado/credentials.env"
PAIR_RE = re.compile(
    r"Cliente:\s*(.+?)\s+CNPJ:\s*(\d{2}\.\d{3}\.\d{3}/\d{4}-\d{2})",
    re.I,
)
JUNK_RE = re.compile(
    r"R\$|Gerar boleto|Fatura|Pedido \d|Será paga|AJUSTE",
    re.I,
)


def nome_cliente(s: str) -> str:
    s = re.sub(r"\s+", " ", s or "").strip()
    half = len(s) // 2
    return s[:half] if len(s) % 2 == 1 and s[:half] == s[half + 1 :] else s


def norm_key(s: str) -> str:
    s = nome_cliente(str(s or ""))
    s = re.sub(r"^N[ºo]\s*\d+\s+", "", s, flags=re.I)
    return re.sub(r"\s+", " ", s).strip().upper().rstrip(".")


def format_cnpj(d: str) -> str | None:
    d = re.sub(r"\D", "", d or "")
    if len(d) != 14:
        return None
    return f"{d[:2]}.{d[2:5]}.{d[5:8]}/{d[8:12]}-{d[12:]}"


def load_cred(key: str) -> str:
    text = CREDS.read_bytes().decode("latin-1")
    for line in text.splitlines():
        if line.startswith(f"{key}="):
            return line.split("=", 1)[1].strip().strip("\"'")
    raise SystemExit(f"faltando {key} em {CREDS}")


def fetch_portal(html_out: Path) -> None:
    user = load_cred("CONECTMAX_FRANQUEADO_USER")
    password = load_cred("CONECTMAX_FRANQUEADO_PASS")
    COOKIE.parent.mkdir(parents=True, exist_ok=True)
    COOKIE.write_text("")
    login = subprocess.run(
        [
            "curl",
            "-sS",
            "-c",
            str(COOKIE),
            "-b",
            str(COOKIE),
            "-X",
            "POST",
            "-H",
            "Content-Type: application/x-www-form-urlencoded",
            "--data-urlencode",
            f"usuario={user}",
            "--data-urlencode",
            f"senha={password}",
            "https://franqueado.conectmax.com/franqueado/ajax/login.php",
        ],
        capture_output=True,
        text=True,
        check=True,
    )
    if "location" not in login.stdout.lower():
        raise SystemExit(f"login falhou: {login.stdout[:200]}")
    subprocess.run(
        [
            "curl",
            "-sS",
            "-b",
            str(COOKIE),
            "-G",
            "--data-urlencode",
            "page=contasareceber",
            "--data-urlencode",
            "data_inicial=01/01/2025",
            "--data-urlencode",
            "data_final=31/12/2026",
            "-o",
            str(html_out),
            "https://franqueado.conectmax.com/franqueado/",
        ],
        check=True,
    )
    print("portal ok", html_out, html_out.stat().st_size)


def parse_portal(html: str) -> dict[str, Counter]:
    name_counts: dict[str, Counter] = defaultdict(Counter)
    for tr in re.finditer(r"<tr[^>]*>(.*?)</tr>", html, re.S | re.I):
        text = unescape(re.sub(r"<[^>]+>", " ", tr.group(1)))
        text = re.sub(r"\s+", " ", text)
        m = PAIR_RE.search(text)
        if not m:
            continue
        nome = nome_cliente(m.group(1).strip())
        if JUNK_RE.search(nome) or len(nome) < 3 or len(nome) > 200:
            continue
        d = re.sub(r"\D", "", m.group(2))
        if len(d) != 14:
            continue
        name_counts[d][nome] += 1
    return name_counts


def ensure(by_cnpj, by_name, cli: str):
    key = norm_key(cli)
    if key in by_name:
        return by_cnpj[by_name[key]]
    synth = f"NOME:{key}"
    by_cnpj[synth] = {
        "cnpj": None,
        "razaoSocial": cli,
        "nomes": {cli},
        "fontes": set(),
        "pedidosQtd": 0,
        "pedidosValor": 0.0,
        "orcamentosQtd": 0,
        "orcamentosValor": 0.0,
        "vendedores": set(),
    }
    by_name[key] = synth
    return by_cnpj[synth]


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--fetch-portal", action="store_true")
    ap.add_argument("--html", type=Path, default=Path("/tmp/cr-all.html"))
    args = ap.parse_args()

    if args.fetch_portal:
        fetch_portal(args.html)

    if not args.html.exists():
        print(
            f"HTML ausente: {args.html}\nRode com --fetch-portal",
            file=sys.stderr,
        )
        sys.exit(1)

    html = args.html.read_bytes().decode("latin-1", errors="ignore")
    name_counts = parse_portal(html)

    by_cnpj = {}
    by_name = {}
    for d, counts in name_counts.items():
        best = counts.most_common(1)[0][0]
        by_cnpj[d] = {
            "cnpj": format_cnpj(d),
            "razaoSocial": best,
            "nomes": set(counts.keys()),
            "fontes": {"portal_contas_receber"},
            "pedidosQtd": 0,
            "pedidosValor": 0.0,
            "orcamentosQtd": 0,
            "orcamentosValor": 0.0,
            "vendedores": set(),
        }
        for n in counts:
            by_name[norm_key(n)] = d

    wb = openpyxl.load_workbook(HERE / "Funil-Pedidos-2025-2026.xlsx", data_only=True)
    for name in wb.sheetnames:
        if not re.fullmatch(r"[A-Z][a-z]{2}\d{4}", name):
            continue
        ws = wb[name]
        header = list(next(ws.iter_rows(min_row=6, max_row=6, values_only=True)))
        for r in ws.iter_rows(min_row=7, values_only=True):
            if not any(v is not None for v in r):
                continue
            row = {h: v for h, v in zip(header, r) if h}
            cli = nome_cliente(str(row.get("Cliente") or ""))
            if not cli:
                continue
            rec = ensure(by_cnpj, by_name, cli)
            rec["fontes"].add("planilha_pedidos")
            rec["pedidosQtd"] += 1
            rec["pedidosValor"] += float(row.get("Total") or 0)
            rec["nomes"].add(cli)
            if row.get("Vendedor"):
                rec["vendedores"].add(str(row.get("Vendedor")))

    wb2 = openpyxl.load_workbook(
        HERE / "Funil-Orcamentos-2025-2026.xlsx", data_only=True
    )
    for name in wb2.sheetnames:
        if not re.fullmatch(r"[A-Z][a-z]{2}\d{4}", name):
            continue
        ws = wb2[name]
        header = list(next(ws.iter_rows(min_row=6, max_row=6, values_only=True)))
        for r in ws.iter_rows(min_row=7, values_only=True):
            if not any(v is not None for v in r):
                continue
            row = {h: v for h, v in zip(header, r) if h}
            raw = re.sub(
                r"^N[ºo]\s*\d+\s+", "", str(row.get("Cliente") or ""), flags=re.I
            )
            cli = nome_cliente(raw)
            if not cli:
                continue
            rec = ensure(by_cnpj, by_name, cli)
            rec["fontes"].add("planilha_orcamentos")
            rec["orcamentosQtd"] += 1
            rec["orcamentosValor"] += float(row.get("Total") or 0)
            rec["nomes"].add(cli)
            if row.get("Vendedor"):
                rec["vendedores"].add(str(row.get("Vendedor")))

    out = []
    for rec in by_cnpj.values():
        nomes = sorted(rec["nomes"], key=lambda n: (-len(n), n))
        uniq = []
        seen = set()
        for n in nomes:
            k = norm_key(n)
            if k in seen:
                continue
            seen.add(k)
            uniq.append(n)
        razao = rec["razaoSocial"]
        if rec["cnpj"] and uniq:
            corp = [n for n in uniq if re.search(r"\b(LTDA|S/?A|S\.A)\b", n, re.I)]
            razao = max(corp or uniq, key=len)
        fantasia = None
        if len(uniq) >= 2:
            short = min(uniq, key=len)
            if norm_key(short) != norm_key(razao):
                fantasia = short
        out.append(
            {
                "cnpj": rec["cnpj"],
                "razaoSocial": razao,
                "nomeFantasia": fantasia,
                "fontes": sorted(rec["fontes"]),
                "pedidosQtd": rec["pedidosQtd"],
                "pedidosValor": round(rec["pedidosValor"], 2),
                "orcamentosQtd": rec["orcamentosQtd"],
                "orcamentosValor": round(rec["orcamentosValor"], 2),
                "vendedores": sorted(rec["vendedores"]),
            }
        )

    out.sort(
        key=lambda x: (-(1 if x["cnpj"] else 0), -x["pedidosValor"], x["razaoSocial"])
    )
    payload = {
        "geradoEm": date.today().isoformat(),
        "fonte": "portal contasareceber (01/01/2025–31/12/2026) + planilhas Pedidos/Orçamentos",
        "total": len(out),
        "comCnpj": sum(1 for x in out if x["cnpj"]),
        "semCnpj": sum(1 for x in out if not x["cnpj"]),
        "empresas": out,
    }
    dest = HERE / "empresas.json"
    dest.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print("ok", dest, "total", payload["total"], "comCnpj", payload["comCnpj"])


if __name__ == "__main__":
    main()
